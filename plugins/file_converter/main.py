# -*- coding: utf-8 -*-
"""
文件格式转换插件
第一期:支持常见图片格式互转
支持格式:PNG ↔ JPG ↔ JPEG ↔ WebP ↔ BMP ↔ ICO ↔ GIF
"""
import os
import uuid
from PIL import Image

# 支持的图片格式映射
SUPPORTED_IMAGE_FORMATS = {
    'png': 'PNG',
    'jpg': 'JPEG',
    'jpeg': 'JPEG',
    'webp': 'WebP',
    'bmp': 'BMP',
    'ico': 'ICO',
    'gif': 'GIF'
}

# 各格式默认保存参数
DEFAULT_SAVE_KWARGS = {
    'JPEG': {'quality': 80, 'optimize': True},
    'WebP': {'quality': 80, 'lossless': False},
    'PNG': {'optimize': True},
    'BMP': {},
    'ICO': {'sizes': [(256, 256), (128, 128), (64, 64), (32, 32), (16, 16)]},
    'GIF': {'optimize': True}
}
# 支持的文档格式映射
SUPPORTED_DOC_FORMATS = {
    'docx': 'Word文档',
    'xlsx': 'Excel表格'
}

def _normalize_markdown_content(md_content):
    """轻量容错处理:自动修正MD偶发格式瑕疵，无感知处理模型输出问题"""
    lines = md_content.splitlines()
    processed_lines = []
    in_table = False
    
    for line in lines:
        stripped = line.strip()
        # 处理表格分隔线缺失问题
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            # 是表头行，自动补全缺失的分隔线
            if not in_table and len(cells) >= 1:
                in_table = True
                processed_lines.append(line)
                separator = '|' + '|'.join(['---'] * len(cells)) + '|'
                processed_lines.append(separator)
                continue
            # 跳过分隔线行
            if all(set(c.strip()) <= {'-', ':'} for c in cells if c.strip()):
                continue
        else:
            in_table = False
        
        # 处理标题跳级问题（自动修正为连续层级）
        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            if level > 1:
                # 查找最近的上一级标题
                has_parent = False
                for prev_line in reversed(processed_lines[-10:]):
                    prev_stripped = prev_line.strip()
                    if prev_stripped.startswith('#' * (level-1) + ' '):
                        has_parent = True
                        break
                if not has_parent:
                    line = '## ' + stripped.lstrip('#').strip()
        
        processed_lines.append(line)
    
    return '\n'.join(processed_lines)

def _validate_output_file(file_path, expected_format, expected_table_count=None):
    """兜底校验:检查输出文件是否有效"""
    try:
        if not os.path.isfile(file_path) or os.path.getsize(file_path) == 0:
            return False
        
        if expected_format == 'docx':
            # 校验docx文件头（PK zip格式）
            with open(file_path, 'rb') as f:
                header = f.read(2)
                return header == b'PK'
        elif expected_format == 'xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True)
            if expected_table_count is not None:
                return len(wb.sheetnames) == expected_table_count
            return True
        return True
    except Exception:
        return False

def _set_docx_table_borders(docx_path):
    """强制给Word文档中所有表格添加0.5磅黑色实线边框，兼容手机端显示"""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        doc = Document(docx_path)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    
                    tcBorders = tcPr.find(qn('w:tcBorders'))
                    if tcBorders is None:
                        tcBorders = OxmlElement('w:tcBorders')
                        tcPr.append(tcBorders)
                    
                    for border_name in ['top', 'left', 'bottom', 'right']:
                        border = tcBorders.find(qn(f'w:{border_name}'))
                        if border is None:
                            border = OxmlElement(f'w:{border_name}')
                            tcBorders.append(border)
                        border.set(qn('w:val'), 'single')
                        border.set(qn('w:sz'), '4')
                        border.set(qn('w:color'), '000000')
                        border.set(qn('w:space'), '0')
        
        doc.save(docx_path)
        return True
    except Exception as e:
        print(f"设置表格边框失败: {str(e)}")
        return False

def _convert_md_to_docx(input_path, output_dir=None):
    """MD转Word文档，使用内置Pandoc绿色版默认模板"""
    try:
        import pypandoc
        from pypandoc.pandoc_download import download_pandoc
        
        # 校验源文件
        if not os.path.isfile(input_path) or not input_path.lower().endswith('.md'):
            return {'success': False, 'input_path': input_path, 'message': '源文件不是有效的Markdown文件'}
        
        # 解析输出路径
        input_dir, input_filename = os.path.split(input_path)
        base_name, _ = os.path.splitext(input_filename)
        if not output_dir:
            output_dir = input_dir
        os.makedirs(output_dir, exist_ok=True)
        output_path = _get_unique_output_path(output_dir, base_name, 'docx')
        
        # 读取并预处理MD内容
        with open(input_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
        processed_content = _normalize_markdown_content(md_content)
        
        # 自动检测Pandoc，缺失则自动下载绿色版
        try:
            pypandoc.get_pandoc_path()
        except OSError:
            download_pandoc()
        
        # 转换，带1次自动重试
        for retry in range(2):
            try:
                pypandoc.convert_text(
                    processed_content,
                    'docx',
                    format='md',
                    outputfile=output_path,
                    extra_args=['--standalone']
                )
                # 校验文件有效性
                if _validate_output_file(output_path, 'docx'):
                    # 强制设置表格实线边框，兼容手机端显示
                    _set_docx_table_borders(output_path)
                    return {
                        'success': True,
                        'input_path': input_path,
                        'output_path': output_path,
                        'file_size': os.path.getsize(output_path),
                        'message': 'Word文档转换成功'
                    }
            except Exception:
                if retry == 1:
                    raise
                continue
        
        return {'success': False, 'input_path': input_path, 'message': '转换失败:生成的文件无效'}
        
    except ImportError:
        return {'success': False, 'message': '依赖缺失:pypandoc-binary未安装，请执行pip install pypandoc-binary==1.13后重试'}
    except Exception as e:
        return {'success': False, 'input_path': input_path, 'message': f'Word转换失败:{str(e)}'}

def _convert_md_to_xlsx(input_path, output_dir=None):
    """提取MD中所有表格转换为Excel，每个表格对应一个Sheet"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter
        
        # 校验源文件
        if not os.path.isfile(input_path) or not input_path.lower().endswith('.md'):
            return {'success': False, 'input_path': input_path, 'message': '源文件不是有效的Markdown文件'}
        
        # 解析输出路径
        input_dir, input_filename = os.path.split(input_path)
        base_name, _ = os.path.splitext(input_filename)
        if not output_dir:
            output_dir = input_dir
        os.makedirs(output_dir, exist_ok=True)
        output_path = _get_unique_output_path(output_dir, base_name, 'xlsx')
        
        # 读取并预处理MD内容
        with open(input_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
        processed_content = _normalize_markdown_content(md_content)
        
        # 提取所有MD表格
        lines = processed_content.splitlines()
        tables = []
        current_table = []
        current_title = f"表格{len(tables)+1}"
        
        for line in lines:
            stripped = line.strip()
            # 检测表格前的标题作为Sheet名
            if stripped.startswith('#'):
                if current_table:
                    tables.append((current_title, current_table))
                    current_table = []
                current_title = stripped.lstrip('#').strip()[:31]  # Excel Sheet名最长31字符
                continue
            # 识别表格行
            if stripped.startswith('|') and stripped.endswith('|'):
                cells = [c.strip() for c in stripped.split('|')[1:-1]]
                # 跳过分隔线行
                if all(set(c) <= {'-', ':'} for c in cells if c):
                    continue
                current_table.append(cells)
            else:
                if current_table:
                    tables.append((current_title, current_table))
                    current_table = []
                    current_title = f"表格{len(tables)+2}"
        # 处理最后一个表格
        if current_table:
            tables.append((current_title, current_table))
        
        if not tables:
            return {'success': False, 'input_path': input_path, 'message': '未在Markdown中识别到任何有效二维表格'}
        
        # 生成Excel
        wb = Workbook()
        # 删除默认Sheet
        wb.remove(wb.active)
        
        for sheet_name, table_data in tables:
            # 处理重复Sheet名
            existing_names = wb.sheetnames
            final_sheet_name = sheet_name
            counter = 1
            while final_sheet_name in existing_names:
                final_sheet_name = f"{sheet_name[:28]}_{counter}"
                counter += 1
            
            ws = wb.create_sheet(title=final_sheet_name)
            # 写入数据
            for row_idx, row_data in enumerate(table_data, 1):
                for col_idx, cell_value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
                    # 表头加粗
                    if row_idx == 1:
                        cell.font = Font(bold=True)
                    cell.alignment = Alignment(vertical='center', wrap_text=True)
            
            # 自适应列宽
            for col_idx in range(1, len(table_data[0]) + 1):
                max_length = 0
                column_letter = get_column_letter(col_idx)
                for row_idx in range(1, len(table_data) + 1):
                    cell_value = str(ws.cell(row=row_idx, column=col_idx).value or '')
                    # 中文按2字符宽度计算
                    length = sum(2 if '\u4e00' <= c <= '\u9fff' else 1 for c in cell_value)
                    if length > max_length:
                        max_length = length
                ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
        
        # 保存并重试验证
        for retry in range(2):
            try:
                wb.save(output_path)
                if _validate_output_file(output_path, 'xlsx', expected_table_count=len(tables)):
                    return {
                        'success': True,
                        'input_path': input_path,
                        'output_path': output_path,
                        'file_size': os.path.getsize(output_path),
                        'table_count': len(tables),
                        'message': f'Excel转换成功，共提取{len(tables)}个表格'
                    }
            except Exception:
                if retry == 1:
                    raise
                continue
        
        return {'success': False, 'input_path': input_path, 'message': '转换失败:生成的文件无效'}
        
    except ImportError:
        return {'success': False, 'message': '依赖缺失:openpyxl未安装，请执行pip install openpyxl==3.1.5后重试'}
    except Exception as e:
        return {'success': False, 'input_path': input_path, 'message': f'Excel转换失败:{str(e)}'}


def init():
    """插件初始化方法，安装/启用/配置修改时自动调用"""
    try:
        import PIL
        all_supported = list(SUPPORTED_IMAGE_FORMATS.keys()) + list(SUPPORTED_DOC_FORMATS.keys())
        return {
            'success': True,
            'message': '文件格式转换插件初始化成功，已支持图片格式互转、MD转Word/Excel',
            'supported_formats': all_supported
        }
    except ImportError:
        return {
            'success': False,
            'message': '依赖缺失:Pillow库未安装，请执行pip install Pillow==10.3.0后重启插件'
        }


def _get_unique_output_path(output_dir, base_name, target_ext):
    """生成唯一的输出路径，同名文件自动加序号后缀，不覆盖原文件"""
    counter = 1
    output_path = os.path.join(output_dir, f"{base_name}.{target_ext}")
    while os.path.exists(output_path):
        output_path = os.path.join(output_dir, f"{base_name}({counter}).{target_ext}")
        counter += 1
    return output_path


def _convert_single_image(input_path, output_format, output_dir=None, quality=80):
    """转换单个图片文件"""
    try:
        # 校验源文件存在
        if not os.path.isfile(input_path):
            return {'success': False, 'input_path': input_path, 'message': '源文件不存在'}

        # 校验目标格式支持
        target_format = SUPPORTED_IMAGE_FORMATS.get(output_format.lower())
        if not target_format:
            return {'success': False, 'input_path': input_path, 'message': f'不支持的目标格式:{output_format}'}

        # 解析输出目录和文件名
        input_dir, input_filename = os.path.split(input_path)
        base_name, _ = os.path.splitext(input_filename)
        if not output_dir:
            output_dir = input_dir
        os.makedirs(output_dir, exist_ok=True)

        # 生成唯一输出路径
        output_path = _get_unique_output_path(output_dir, base_name, output_format.lower())

        # 打开图片并处理模式
        with Image.open(input_path) as img:
            # 处理透明通道:转JPG/BMP等不支持透明的格式时，添加白色背景
            if target_format in ['JPEG', 'BMP'] and img.mode in ('RGBA', 'LA', 'P'):
                background = Image.new('RGB', img.size, (255, 255, 255))
                if img.mode == 'P':
                    img = img.convert('RGBA')
                background.paste(img, mask=img.split()[-1] if img.mode in ('RGBA', 'LA') else None)
                img = background
            elif target_format == 'ICO' and img.mode not in ('RGBA', 'RGB'):
                img = img.convert('RGBA')
            elif target_format == 'GIF' and img.mode not in ('P', 'RGB', 'RGBA'):
                img = img.convert('P', palette=Image.ADAPTIVE)

            # 组装保存参数
            save_kwargs = DEFAULT_SAVE_KWARGS.get(target_format, {}).copy()
            if target_format in ['JPEG', 'WebP'] and 1 <= quality <= 100:
                save_kwargs['quality'] = quality

            # 保存文件
            img.save(output_path, format=target_format, **save_kwargs)

        # 获取文件大小
        file_size = os.path.getsize(output_path)
        return {
            'success': True,
            'input_path': input_path,
            'output_path': output_path,
            'file_size': file_size,
            'message': '转换成功'
        }

    except Exception as e:
        return {
            'success': False,
            'input_path': input_path,
            'message': f'转换失败:{str(e)}'
        }


def run(params):
    """插件主入口"""
    input_path = params.get('input_path', '').strip()
    output_format = params.get('output_format', '').strip()
    output_dir = params.get('output_dir', '').strip()
    quality = params.get('quality', 80)

    # 参数校验
    if not input_path:
        return {'success': False, 'message': '参数错误:input_path不能为空'}
    if not output_format:
        return {'success': False, 'message': '参数错误:output_format不能为空'}
    
    output_format_lower = output_format.lower()
    all_supported_formats = list(SUPPORTED_IMAGE_FORMATS.keys()) + list(SUPPORTED_DOC_FORMATS.keys())
    if output_format_lower not in all_supported_formats:
        return {'success': False, 'message': f'不支持的目标格式，当前支持:{", ".join(all_supported_formats)}'}

    # 文档格式转换分支（仅支持单MD文件）
    if output_format_lower in SUPPORTED_DOC_FORMATS:
        if not os.path.isfile(input_path):
            return {'success': False, 'message': '文档转换仅支持单个Markdown文件，不支持目录批量转换'}
        if output_format_lower == 'docx':
            result = _convert_md_to_docx(input_path, output_dir or None)
        else:
            result = _convert_md_to_xlsx(input_path, output_dir or None)
        return {
            'success': result['success'],
            'total': 1,
            'success_count': 1 if result['success'] else 0,
            'fail_count': 0 if result['success'] else 1,
            'results': [result],
            'message': result['message']
        }

    results = []
    success_count = 0
    fail_count = 0

    # 处理单文件
    if os.path.isfile(input_path):
        result = _convert_single_image(input_path, output_format, output_dir or None, quality)
        results.append(result)
        if result['success']:
            success_count += 1
        else:
            fail_count += 1

    # 处理目录批量转换
    elif os.path.isdir(input_path):
        for filename in os.listdir(input_path):
            file_path = os.path.join(input_path, filename)
            if not os.path.isfile(file_path):
                continue
            # 检查文件扩展名是否为支持的图片格式
            ext = os.path.splitext(filename)[1].lower().lstrip('.')
            if ext in SUPPORTED_IMAGE_FORMATS:
                result = _convert_single_image(file_path, output_format, output_dir or None, quality)
                results.append(result)
                if result['success']:
                    success_count += 1
                else:
                    fail_count += 1

    else:
        return {'success': False, 'message': f'输入路径不存在:{input_path}'}

    return {
        'success': True,
        'total': len(results),
        'success_count': success_count,
        'fail_count': fail_count,
        'results': results,
        'message': f'转换完成:成功{success_count}个，失败{fail_count}个'
    }


def uninstall():
    """插件卸载时调用"""
    return {'success': True, 'message': '文件格式转换插件卸载成功'}